"""Public file tools. Path/index helpers are private."""

from __future__ import annotations

import csv
import io
import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import config
import logger
import paging
import store

try:
    from pypdf import PdfReader, PdfWriter
except ImportError:  # pragma: no cover
    PdfReader = None  # type: ignore
    PdfWriter = None  # type: ignore

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None  # type: ignore

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore
    load_workbook = None  # type: ignore

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas as pdf_canvas
except ImportError:  # pragma: no cover
    letter = None  # type: ignore
    pdf_canvas = None  # type: ignore


def search_dir(
    directory: str,
    name: str | None = None,
    type: str | None = None,
    recursive: bool = False,
    agent_id: str | None = None,
) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    root = verifyPath(directory)
    if root is None:
        return _done("search_dir", run, agent_id, {"ok": False, "error": "path_denied"})
    hits: list[dict[str, str]] = []
    iterator = root.rglob("*") if recursive else root.iterdir()
    try:
        for item in iterator:
            if not item.is_file():
                continue
            if not _allowed_file(item):
                continue
            if name and name.lower() not in item.name.lower():
                continue
            if type and item.suffix.lower() != _dot(type):
                continue
            hits.append({"path": str(item), "name": item.name})
            if len(hits) >= config.SEARCH_CAP:
                break
    except OSError as exc:
        return _done("search_dir", run, agent_id, {"ok": False, "error": type(exc).__name__})
    return _done("search_dir", run, agent_id, {"ok": True, "hits": hits, "count": len(hits)})


def read_file(
    directory: str | None = None,
    name: str | None = None,
    file_path: str | None = None,
    start: int = 1,
    end: int = 1,
    agent_id: str | None = None,
) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    target = _join_or_path(directory, name, file_path)
    if target is None or not _allowed_file(target):
        return _done("read_file", run, agent_id, {"ok": False, "error": "path_denied"})
    if not target.is_file():
        target = _read_with_suffix(directory, name, file_path) or target
    if not target.is_file():
        return _done("read_file", run, agent_id, {"ok": False, "error": "not_found"})
    try:
        payload = _read_dispatch(target, _as_page(start, 1), _as_page(end, 1))
    except Exception as exc:
        return _done("read_file", run, agent_id, {"ok": False, "error": type(exc).__name__})
    payload["ok"] = True
    payload["path"] = str(target)
    return _done("read_file", run, agent_id, payload, summary=f"path={target.name}")


def write_doc(
    directory: str | None = None,
    name: str = "",
    type: str = "txt",
    content: str = "",
    structure: dict[str, Any] | None = None,
    agent_id: str | None = None,
) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    if not verifyType(type) or type == "pdf":
        return _done("write_doc", run, agent_id, {"ok": False, "error": "type_denied"})
    dest = _keep_dest(run, _join_write(directory, name, type))
    if dest is None:
        return _done("write_doc", run, agent_id, {"ok": False, "error": "path_denied"})
    if type in {"txt", "md"} and _placeholder_text(content) and not structure:
        return _done("write_doc", run, agent_id, {"ok": False, "error": "empty_content"})
    gate = _write_gate(dest)
    if not gate["ok"]:
        return _done("write_doc", run, agent_id, gate)
    try:
        data = _render_doc(type, content, structure or {})
        if not paging.char_chunk_ok(content if isinstance(content, str) else json.dumps(structure or {})):
            return _done("write_doc", run, agent_id, {"ok": False, "error": "chunk_too_large"})
        _atomic_write(dest, data)
        _index_write(dest, agent_id, run, "write_doc", gate.get("stale", False))
    except Exception as exc:
        return _done("write_doc", run, agent_id, {"ok": False, "error": type(exc).__name__})
    return _done("write_doc", run, agent_id, {"ok": True, "path": str(dest)})


def write_pdf(
    directory: str | None = None,
    name: str = "",
    content: str = "",
    structure: dict[str, Any] | None = None,
    agent_id: str | None = None,
) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    dest = _keep_dest(run, _join_write(directory, name, "pdf"))
    if dest is None:
        return _done("write_pdf", run, agent_id, {"ok": False, "error": "path_denied"})
    gate = _write_gate(dest)
    if not gate["ok"]:
        return _done("write_pdf", run, agent_id, gate)
    if not paging.char_chunk_ok(content):
        return _done("write_pdf", run, agent_id, {"ok": False, "error": "chunk_too_large"})
    try:
        _write_pdf_bytes(dest, content, structure or {})
        _index_write(dest, agent_id, run, "write_pdf", gate.get("stale", False))
    except Exception as exc:
        return _done("write_pdf", run, agent_id, {"ok": False, "error": type(exc).__name__})
    return _done("write_pdf", run, agent_id, {"ok": True, "path": str(dest)})


def write_json(
    directory: str | None = None,
    name: str = "",
    content: Any = None,
    agent_id: str | None = None,
) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    dest = _keep_dest(run, _join_write(directory, name, "json"))
    if dest is None:
        return _done("write_json", run, agent_id, {"ok": False, "error": "path_denied"})
    gate = _write_gate(dest)
    if not gate["ok"]:
        return _done("write_json", run, agent_id, gate)
    try:
        parsed = json.loads(content) if isinstance(content, str) else content
        mismatch = _json_shape_mismatch(dest, parsed)
        if mismatch:
            return _done("write_json", run, agent_id, {"ok": False, "error": mismatch})
        raw = json.dumps(parsed, indent=2)
        if not paging.char_chunk_ok(raw):
            return _done("write_json", run, agent_id, {"ok": False, "error": "chunk_too_large"})
        _atomic_write(dest, raw.encode("utf-8"))
        _index_write(dest, agent_id, run, "write_json", gate.get("stale", False))
    except (json.JSONDecodeError, TypeError, ValueError):
        return _done("write_json", run, agent_id, {"ok": False, "error": "invalid_json"})
    except Exception as exc:
        return _done("write_json", run, agent_id, {"ok": False, "error": type(exc).__name__})
    return _done("write_json", run, agent_id, {"ok": True, "path": str(dest)})


def append_file(
    directory: str | None = None,
    name: str = "",
    type: str = "txt",
    content: Any = "",
    agent_id: str | None = None,
) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    dest = _keep_dest(run, _join_write(directory, name, type))
    if dest is None:
        return _done("append_file", run, agent_id, {"ok": False, "error": "path_denied"})
    index = _load_index()
    key = str(dest)
    if key not in index:
        return _done("append_file", run, agent_id, {"ok": False, "error": "not_in_index"})
    if not dest.is_file():
        return _done("append_file", run, agent_id, {"ok": False, "error": "missing_file"})
    if type in {"doc", "docx", "xlsx"}:
        return _done("append_file", run, agent_id, {"ok": False, "error": "no_append"})
    try:
        if type == "txt":
            if not isinstance(content, str) or not paging.char_chunk_ok(content):
                return _done("append_file", run, agent_id, {"ok": False, "error": "chunk_too_large"})
            with dest.open("a", encoding="utf-8") as fh:
                fh.write(content)
        elif type == "csv":
            rows = _as_rows(content)
            if not paging.row_chunk_ok(rows):
                return _done("append_file", run, agent_id, {"ok": False, "error": "chunk_too_large"})
            with dest.open("a", encoding="utf-8", newline="") as fh:
                csv.writer(fh).writerows(rows)
        elif type == "json":
            _append_json(dest, content)
        elif type == "pdf":
            _append_pdf(dest, content)
        else:
            return _done("append_file", run, agent_id, {"ok": False, "error": "type_denied"})
        index[key]["last_writer"] = agent_id
        index[key]["last_write_at"] = _now()
        _save_index(index)
    except Exception as exc:
        return _done("append_file", run, agent_id, {"ok": False, "error": type(exc).__name__})
    return _done("append_file", run, agent_id, {"ok": True, "path": str(dest)})


def mkdir(directory: str | None = None, name: str = "", agent_id: str | None = None) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    directory, name = _mkdir_parts(directory, name)
    parent = _verify_write_dir(directory)
    if parent is None:
        return _done("mkdir", run, agent_id, {"ok": False, "error": "path_denied"})
    if Path(name).name != name or not name or name in {".", ".."}:
        return _done("mkdir", run, agent_id, {"ok": False, "error": "bad_name"})
    target = (parent / name).resolve()
    target = _keep_dest(run, target)
    if target is None or not _writable(target) or _blocked(target):
        return _done("mkdir", run, agent_id, {"ok": False, "error": "path_denied"})
    if run is None or getattr(run, "keep_path", None) is None or getattr(run, "level", None) != "large":
        if not _is_under(target, parent):
            return _done("mkdir", run, agent_id, {"ok": False, "error": "path_denied"})
    try:
        target.mkdir(parents=False, exist_ok=True)
    except Exception as exc:
        return _done("mkdir", run, agent_id, {"ok": False, "error": type(exc).__name__})
    return _done("mkdir", run, agent_id, {"ok": True, "path": str(target)})


def keep_copy(dest: Path, data: bytes, agent_id: str | None = None) -> dict[str, Any]:
    run, agent_id = _run_agent(agent_id)
    target = Path(dest).resolve()
    keep = None if run is None else getattr(run, "keep_path", None)
    if keep is not None:
        root = Path(keep).resolve()
        if target != root and not _is_under(target, root):
            return _done("keep_copy", run, agent_id, {"ok": False, "error": "path_denied"})
    if not _writable(target) or _blocked(target):
        return _done("keep_copy", run, agent_id, {"ok": False, "error": "path_denied"})
    gate = _write_gate(target)
    if not gate["ok"]:
        return _done("keep_copy", run, agent_id, gate)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        _atomic_write(target, data)
        _index_write(target, agent_id, run, "keep_copy", gate.get("stale", False))
    except Exception as exc:
        return _done("keep_copy", run, agent_id, {"ok": False, "error": type(exc).__name__})
    return _done("keep_copy", run, agent_id, {"ok": True, "path": str(target)})


def verifyPath(directory: str | Path) -> Path | None:
    raw = Path(directory)
    if ".." in raw.parts:
        return None
    try:
        resolved = raw.resolve()
    except OSError:
        return None
    if _blocked(resolved):
        return None
    if not any(_is_under(resolved, root) for root in config.ALLOWED_PATHS):
        return None
    return resolved


def verifyType(doc_type: str) -> bool:
    return doc_type in config.ALLOWED_DOC_TYPES


def _done(tool: str, run: Any, agent_id: str | None, result: dict[str, Any], summary: str = "") -> dict[str, Any]:
    return logger.finish(
        event="tool",
        tool=tool,
        run_id=None if run is None else run.run_id,
        agent_name=agent_id,
        result=result,
        level=None if run is None else run.level,
        stakes=None if run is None else run.stakes,
        mail_cal=None if run is None else run.mail_cal,
        summary=summary,
    )


def _run_agent(agent_id: str | None) -> tuple[Any, str | None]:
    if not agent_id:
        return None, None
    return store.run_for_agent(agent_id), agent_id


def _join_or_path(directory: str | None, name: str | None, file_path: str | None) -> Path | None:
    if file_path:
        target = Path(file_path)
        if ".." in target.parts:
            return None
        try:
            resolved = target.resolve()
        except OSError:
            return None
        parent = verifyPath(resolved.parent)
        if parent is None:
            return None
        return resolved if _is_under(resolved, parent) else None
    if directory is None or name is None:
        return None
    return _join_read(directory, name)


def _join_read(directory: str, name: str) -> Path | None:
    if Path(name).name != name or not name:
        return None
    raw = Path(directory)
    if ".." in raw.parts:
        return None
    candidates: list[Path] = []
    if raw.is_absolute():
        candidates.append((raw / name).resolve())
    else:
        parts = list(raw.parts)
        for root in config.ALLOWED_PATHS:
            rest = parts[1:] if parts and parts[0].lower() == root.name.lower() else parts
            base = root.joinpath(*rest) if rest else root
            candidates.append((base / name).resolve())
    seen: set[Path] = set()
    for dest in candidates:
        if dest in seen:
            continue
        seen.add(dest)
        parent = verifyPath(dest.parent)
        if parent is None:
            continue
        if _is_under(dest, parent) and not _blocked(dest):
            return dest
    return None


def _join_write(directory: str | None, name: str, doc_type: str, add_ext: bool = True) -> Path | None:
    parent = _verify_write_dir(directory)
    if parent is None:
        return None
    if Path(name).name != name or not name:
        return None
    stem = Path(name).stem if add_ext else name
    if add_ext:
        dest = (parent / f"{stem}{_dot(doc_type)}").resolve()
    else:
        dest = (parent / name).resolve()
    if not _is_under(dest, parent) or not _writable(dest) or _blocked(dest):
        return None
    return dest


def _keep_dest(run: Any, dest: Path | None) -> Path | None:
    if dest is None:
        return None
    keep = None if run is None else getattr(run, "keep_path", None)
    if run is None or getattr(run, "level", None) != "large" or keep is None:
        return dest
    keep_root = Path(keep).resolve()
    resolved = dest.resolve()
    try:
        resolved.relative_to(keep_root)
        return resolved
    except ValueError:
        remapped = (keep_root / dest.name).resolve()
        if not _writable(remapped) or _blocked(remapped):
            return None
        return remapped


def _verify_write_dir(directory: str | None) -> Path | None:
    roots = [p.resolve() for p in config.WRITE_PATHS] or [config.OUTPUTS.resolve()]
    default = roots[0]
    if not directory:
        resolved = verifyPath(default)
        return resolved if resolved is not None and _writable(resolved) else None
    raw = Path(directory)
    if ".." in raw.parts:
        return None
    if not raw.is_absolute():
        parts = list(raw.parts)
        if parts and parts[0].lower() == default.name.lower():
            parts = parts[1:]
        raw = default.joinpath(*parts) if parts else default
    resolved = verifyPath(raw)
    if resolved is None or not _writable(resolved):
        return None
    return resolved


def _dot(doc_type: str) -> str:
    t = doc_type if doc_type.startswith(".") else f".{doc_type}"
    return t.lower()


def _as_page(value: Any, default: int) -> int:
    try:
        n = int(value)
    except (TypeError, ValueError):
        return default
    return n if n > 0 else default


def _placeholder_text(content: Any) -> bool:
    text = str(content or "").strip()
    if not text:
        return True
    lowered = text.lower()
    if lowered in {"...", "…", "n/a", "failed", "none", "null", "placeholder"}:
        return True
    if lowered.startswith("[") and "from" in lowered and "]" in lowered:
        return True
    return False


def _json_shape_mismatch(dest: Path, parsed: Any) -> str | None:
    if not dest.is_file():
        return None
    try:
        existing = json.loads(dest.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, UnicodeDecodeError):
        return None
    if isinstance(existing, dict) and isinstance(parsed, list):
        return "json_shape"
    if isinstance(existing, list) and isinstance(parsed, dict):
        return "json_shape"
    return None


def _mkdir_parts(directory: str | None, name: str) -> tuple[str | None, str]:
    if str(name or "").strip():
        return directory, name
    raw = Path(str(directory or "").strip())
    if not directory or ".." in raw.parts or raw.is_absolute() or len(raw.parts) != 1:
        return directory, name
    if raw.name.lower() in {"outputs", "output"}:
        return directory, name
    return None, raw.name


def _read_with_suffix(directory: str | None, name: str | None, file_path: str | None) -> Path | None:
    if file_path or not name or "." in Path(name).name:
        return None
    for ext in (".txt", ".md", ".json", ".pdf", ".csv", ".docx", ".doc"):
        alt = _join_or_path(directory, f"{name}{ext}", None)
        if alt is not None and alt.is_file() and _allowed_file(alt):
            return alt
    return None


def _allowed_file(path: Path) -> bool:
    parent = verifyPath(path.parent)
    return parent is not None and _is_under(path.resolve(), parent) and not _blocked(path)


def _write_locked(path: Path) -> bool:
    resolved = path.resolve()
    for raw in config.INPUT_PATHS:
        try:
            root = raw.resolve()
        except OSError:
            root = raw
        if resolved == root or _is_under(resolved, root):
            return True
    return False


def _writable(path: Path) -> bool:
    if _write_locked(path):
        return False
    resolved = path.resolve()
    for raw in config.WRITE_PATHS:
        try:
            root = raw.resolve()
        except OSError:
            root = raw
        if resolved == root or _is_under(resolved, root):
            return True
    return False


def _blocked(path: Path) -> bool:
    resolved = path.resolve()
    for blocked in config.BLOCKED_PATHS:
        try:
            b = blocked.resolve()
        except OSError:
            b = blocked
        if resolved == b or _is_under(resolved, b):
            if _is_open_root(resolved) and _is_repo_block(b):
                continue
            return True
    if resolved.name.startswith("token_google_") or resolved.name.startswith("token_ms_"):
        return True
    return False


def _is_under(child: Path, parent: Path) -> bool:
    try:
        child.resolve().relative_to(parent.resolve())
        return True
    except (ValueError, OSError):
        return False


def _is_open_root(path: Path) -> bool:
    return any(path == root or _is_under(path, root) for root in config.ALLOWED_PATHS)


def _is_repo_block(blocked: Path) -> bool:
    if blocked == config.REPO_ROOT:
        return True
    return _is_under(config.WORKSPACE, blocked) or _is_under(config.OUTPUTS, blocked)


def _write_gate(dest: Path) -> dict[str, Any]:
    index = _load_index()
    key = str(dest)
    exists = dest.exists()
    indexed = key in index
    if not exists and not indexed:
        return {"ok": True, "stale": False}
    if not exists and indexed:
        return {"ok": True, "stale": True}
    if exists and indexed:
        return {"ok": True, "stale": False}
    return {"ok": False, "error": "original_file"}


def _index_write(dest: Path, agent_id: str | None, run: Any, tool: str, stale: bool) -> None:
    index = _load_index()
    key = str(dest)
    now = _now()
    run_id = None if run is None else run.run_id
    if key not in index:
        index[key] = {
            "path": key,
            "created_by": agent_id,
            "created_at": now,
            "run_id": run_id,
            "tool": tool,
            "last_writer": agent_id,
            "last_write_at": now,
        }
    else:
        if not stale:
            index[key]["last_writer"] = agent_id
            index[key]["last_write_at"] = now
        else:
            index[key]["last_writer"] = agent_id
            index[key]["last_write_at"] = now
    _save_index(index)


def _load_index() -> dict[str, Any]:
    if not config.CREATED_INDEX.is_file():
        return {}
    try:
        return json.loads(config.CREATED_INDEX.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def pdf_paths_for_run(run_id: str, keep_path: Path | None = None) -> list[Path]:
    """PDFs this run wrote under outputs/. Agents never call this."""
    rid = str(run_id or "").strip()
    root = config.OUTPUTS.resolve()
    seen: set[str] = set()
    out: list[Path] = []

    def add(path: Path) -> None:
        try:
            dest = path.resolve()
        except OSError:
            return
        if not dest.is_file() or dest.suffix.lower() != ".pdf":
            return
        try:
            dest.relative_to(root)
        except ValueError:
            return
        key = str(dest)
        if key in seen:
            return
        seen.add(key)
        out.append(dest)

    if rid:
        for key, meta in _load_index().items():
            if not isinstance(meta, dict):
                continue
            if str(meta.get("run_id") or "") != rid:
                continue
            add(Path(str(meta.get("path") or key)))
    if keep_path is not None:
        folder = Path(keep_path)
        if folder.is_dir():
            for path in sorted(folder.rglob("*.pdf")):
                add(path)
    return out


def _save_index(index: dict[str, Any]) -> None:
    config.CREATED_INDEX.parent.mkdir(parents=True, exist_ok=True)
    _atomic_write(config.CREATED_INDEX, json.dumps(index, indent=2).encode("utf-8"))


def _atomic_write(dest: Path, data: bytes) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".tmp-", dir=str(dest.parent))
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        os.replace(tmp, dest)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _read_dispatch(path: Path, start: int, end: int) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return readPdf(path, start, end)
    if suffix == ".csv":
        return readCsv(path, start, end)
    if suffix == ".xlsx":
        return readXlsx(path, start, end)
    if suffix == ".docx":
        return readDocx(path, start, end)
    if suffix == ".json":
        text = path.read_text(encoding="utf-8")
        chunk, s, e, total = paging.slice_chars(text, start, end)
        return {"text": chunk, "start": s, "end": e, "total_pages": total}
    if suffix == ".doc":
        text = _read_doc(path)
        chunk, s, e, total = paging.slice_chars(text, start, end)
        return {"text": chunk, "start": s, "end": e, "total_pages": total}
    text = path.read_text(encoding="utf-8", errors="replace")
    chunk, s, e, total = paging.slice_chars(text, start, end)
    return {"text": chunk, "start": s, "end": e, "total_pages": total}


def readPdf(path: Path, start: int, end: int) -> dict[str, Any]:
    if PdfReader is None:
        raise RuntimeError("pypdf_missing")
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    total = max(1, len(pages))
    s, e = paging.clamp_window(start, end)
    e = min(e, total)
    text = "\n".join(pages[s - 1 : e])
    return {"text": text, "start": s, "end": e, "total_pages": total}


def readCsv(path: Path, start: int, end: int) -> dict[str, Any]:
    with path.open("r", encoding="utf-8", newline="") as fh:
        rows = list(csv.reader(fh))
    chunk, s, e, total = paging.slice_rows(rows, start, end)
    buf = io.StringIO()
    csv.writer(buf).writerows(chunk)
    return {"text": buf.getvalue(), "start": s, "end": e, "total_pages": total}


def readXlsx(path: Path, start: int, end: int) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("openpyxl_missing")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    chunk, s, e, total = paging.slice_rows(rows, start, end)
    buf = io.StringIO()
    csv.writer(buf).writerows(chunk)
    return {"text": buf.getvalue(), "start": s, "end": e, "total_pages": total}


def readDocx(path: Path, start: int, end: int) -> dict[str, Any]:
    if Document is None:
        raise RuntimeError("docx_missing")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    chunk, s, e, total = paging.slice_chars(text, start, end)
    return {"text": chunk, "start": s, "end": e, "total_pages": total}


def _read_doc(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith(b"{\\rtf"):
        text = raw.decode("latin-1", errors="replace")
        return "".join(ch if ch.isprintable() or ch in "\n\r\t" else " " for ch in text)
    return raw.decode("utf-8", errors="replace")


def _render_doc(doc_type: str, content: str, structure: dict[str, Any]) -> bytes:
    if doc_type in {"txt", "md"}:
        return content.encode("utf-8")
    if doc_type == "csv":
        rows = structure.get("rows") or _as_rows(content)
        buf = io.StringIO()
        csv.writer(buf).writerows(rows)
        return buf.getvalue().encode("utf-8")
    if doc_type == "xlsx":
        return _xlsx_bytes(structure.get("rows") or _as_rows(content))
    if doc_type == "docx":
        return _docx_bytes(content, structure)
    if doc_type == "doc":
        body = content or "\n".join(structure.get("paragraphs") or [])
        escaped = body.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}")
        rtf = r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Calibri;}}\f0\fs22 " + escaped.replace("\n", r"\par ") + "}"
        return rtf.encode("utf-8")
    raise ValueError("type_denied")


def _xlsx_bytes(rows: list[list[Any]]) -> bytes:
    if Workbook is None:
        raise RuntimeError("openpyxl_missing")
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(content: str, structure: dict[str, Any]) -> bytes:
    if Document is None:
        raise RuntimeError("docx_missing")
    doc = Document()
    if structure.get("title"):
        doc.add_heading(str(structure["title"]), 0)
    paragraphs = structure.get("paragraphs") or ([content] if content else [])
    for para in paragraphs:
        doc.add_paragraph(str(para))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _write_pdf_bytes(dest: Path, content: str, structure: dict[str, Any]) -> None:
    if pdf_canvas is None or letter is None:
        raise RuntimeError("reportlab_missing")
    lines = []
    if structure.get("title"):
        lines.append(str(structure["title"]))
    lines.extend(structure.get("paragraphs") or [])
    if content:
        lines.extend(content.splitlines() or [content])
    fd, tmp = tempfile.mkstemp(prefix=".tmp-", suffix=".pdf", dir=str(dest.parent))
    os.close(fd)
    try:
        c = pdf_canvas.Canvas(tmp, pagesize=letter)
        width, height = letter
        y = height - 54
        for line in lines or [""]:
            for chunk in _wrap(str(line), 90):
                if y < 54:
                    c.showPage()
                    y = height - 54
                c.drawString(54, y, chunk)
                y -= 14
        c.save()
        os.replace(tmp, dest)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _wrap(text: str, width: int) -> list[str]:
    if not text:
        return [""]
    words = text.split(" ")
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if len(trial) <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def _as_rows(content: Any) -> list[list[Any]]:
    if isinstance(content, list):
        return [row if isinstance(row, list) else [row] for row in content]
    if not isinstance(content, str):
        content = str(content)
    return list(csv.reader(io.StringIO(content)))


def _append_json(dest: Path, content: Any) -> None:
    existing = json.loads(dest.read_text(encoding="utf-8"))
    incoming = json.loads(content) if isinstance(content, str) else content
    if isinstance(existing, list) and isinstance(incoming, list):
        existing.extend(incoming)
        merged = existing
    elif isinstance(existing, dict) and isinstance(incoming, dict):
        existing.update(incoming)
        merged = existing
    elif isinstance(existing, list):
        existing.append(incoming)
        merged = existing
    else:
        raise ValueError("json_extend")
    raw = json.dumps(merged, indent=2)
    if not paging.char_chunk_ok(raw) and len(json.dumps(incoming)) > paging.CHAR_PAGE * paging.MAX_PAGE_SPAN:
        raise ValueError("chunk_too_large")
    _atomic_write(dest, raw.encode("utf-8"))


def _append_pdf(dest: Path, content: Any) -> None:
    if PdfReader is None or PdfWriter is None:
        raise RuntimeError("pypdf_missing")
    writer = PdfWriter()
    writer.append(str(dest))
    extra = Path(content) if isinstance(content, str) and Path(str(content)).is_file() else None
    if extra:
        writer.append(str(extra))
    else:
        tmp = dest.with_suffix(".append-tmp.pdf")
        _write_pdf_bytes(tmp, str(content), {})
        writer.append(str(tmp))
        tmp.unlink(missing_ok=True)
    fd, tmp_out = tempfile.mkstemp(prefix=".tmp-", suffix=".pdf", dir=str(dest.parent))
    os.close(fd)
    try:
        with open(tmp_out, "wb") as fh:
            writer.write(fh)
        os.replace(tmp_out, dest)
    except Exception:
        try:
            os.unlink(tmp_out)
        except OSError:
            pass
        raise
